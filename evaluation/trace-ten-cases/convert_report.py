import json
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).parent
stem = sys.argv[1] if len(sys.argv) > 1 else 'IRIS_TRACE_Accuracy_Evaluation'
tokens = json.loads((ROOT / (stem + '_tokens.json')).read_text(encoding='utf-8'))
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(11.7), Inches(8.3)
s.top_margin = s.bottom_margin = Inches(.65)
s.left_margin = s.right_margin = Inches(.7)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(25)
for name, size in [('Heading 1', 17), ('Heading 2', 13), ('Heading 3', 11)]:
    doc.styles[name].font.size = Pt(size)

expected = []
actual = []
def inline(p, items, bold=False, italic=False):
    for t in items:
        kind = t['type']
        if kind in ('strong', 'em', 'link', 'escape') and t.get('tokens'):
            inline(p, t['tokens'], bold or kind == 'strong', italic or kind == 'em')
            if kind == 'link' and t.get('href') and t['href'] != t['text']:
                p.add_run(' (' + t['href'] + ')')
        elif kind == 'br':
            p.add_run().add_break()
        else:
            text = t.get('text', t.get('raw', ''))
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            if kind == 'codespan':
                r.font.name = 'Consolas'
                r.font.size = Pt(9)
            expected.append(text)
            actual.append(r.text)

def build(items):
    for t in items:
        typ = t['type']
        if typ == 'space':
            continue
        if typ == 'heading':
            p = doc.add_paragraph(style='Title' if t['depth'] == 1 else f'Heading {min(t["depth"]-1,3)}')
            if t['depth'] == 2 and t['text'].startswith('Case '):
                p.paragraph_format.page_break_before = True
            inline(p, t['tokens'])
        elif typ in ('paragraph', 'text'):
            inline(doc.add_paragraph(), t.get('tokens', [{'type': 'text', 'text': t['text']}]))
        elif typ == 'list':
            for i, item in enumerate(t['items']):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(.18)
                p.add_run(f'{int(t.get("start") or 1)+i}. ' if t['ordered'] else '\u2022 ')
                for child in item['tokens']:
                    inline(p, child.get('tokens', [{'type': 'text', 'text': child.get('text','')}]))
        elif typ == 'table':
            n = len(t['header'])
            tab = doc.add_table(rows=1, cols=n)
            tab.autofit = False
            weights = {8:[.45,1.65,1.5,.5, .85,.7,.7,.55], 5:[.55,4.4,1.2,.75,.85],4:[2.1,1,1,5],3:[1.3,1.8,5]}.get(n,[1]*n)
            headers = [cell['text'] for cell in t['header']]
            if stem in ('IRIS_Beginner_Reference_Workbook', 'IRIS_Solo_Calibration_Workbook'):
                if n == 3:
                    weights = [1.1, 6.8, 2.4] if headers[0] == 'ID' else [2, 5.6, 2.7]
                elif n == 4:
                    weights = [.8, 2.3, 3.6, 3.6] if headers[0] == 'ID' else [3.7, 2.2, 2.2, 2.2]
                elif n == 2:
                    weights = [2, 8.3]
                if headers[0] == 'Claim or component ID':
                    weights = [1.2, 2.3, 6.8]
            widths = [10.3*x/sum(weights) for x in weights]
            for col, width in zip(tab.columns, widths): col.width = Inches(width)
            for ri, row in enumerate([t['header']] + t['rows']):
                cells = tab.rows[0].cells if ri == 0 else tab.add_row().cells
                if ri == 0:
                    flag = OxmlElement('w:tblHeader'); tab.rows[0]._tr.get_or_add_trPr().append(flag)
                for ci, celltoken in enumerate(row):
                    cell = cells[ci]; cell.width = Inches(widths[ci])
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(4)
                    inline(p, celltoken['tokens'], ri == 0)
                    for r in p.runs: r.font.size = Pt(9)
                    tcpr = cell._tc.get_or_add_tcPr()
                    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E7E9EC' if ri == 0 else 'FFFFFF'); tcpr.append(shade)
            borders = OxmlElement('w:tblBorders')
            for side in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            tab._tbl.tblPr.append(borders)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif typ == 'blockquote': build(t['tokens'])
        elif typ == 'hr': doc.add_paragraph()
        elif typ == 'code':
            p=doc.add_paragraph(); inline(p,[{'type':'codespan','text':t['text']}])
        else: raise ValueError(f'Unhandled block {typ}')

build(tokens)
footer = s.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('IRIS TRACE Evaluation | ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
out = ROOT / (stem + '.docx')
doc.save(out)
assert expected == actual
print(f'Created {out}; {len(doc.tables)} tables; {len(expected)} preserved text runs')
